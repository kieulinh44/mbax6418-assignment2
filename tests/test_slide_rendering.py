"""Slide rendering regression tests, no external model calls."""

# Heavy model stack required (bm25s / chromadb / sentence-transformers).
# Skip on light test environments (CI with requirements-dev.txt) instead of failing.
import pytest
pytest.importorskip("bm25s")
pytest.importorskip("sentence_transformers")
pytest.importorskip("chromadb")

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from app import ingest


class SlideRenderingTests(unittest.TestCase):
    def test_finds_soffice_on_path_as_a_complete_filename(self):
        with tempfile.TemporaryDirectory() as td:
            executable = Path(td) / 'soffice'
            executable.touch()
            with patch.object(ingest, '_SOFFICE', None), patch.object(ingest.shutil, 'which', return_value=str(executable)):
                self.assertEqual(ingest._find_soffice(), str(executable))
    def test_mismatched_render_count_never_labels_wrong_slide(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            deck=root/'two-slides.pptx'
            prs=Presentation()
            prs.slides.add_slide(prs.slide_layouts[0])
            prs.slides.add_slide(prs.slide_layouts[0])
            prs.save(str(deck))
            with patch.object(ingest.config,'DATA_INDEX',td), patch.object(ingest,'_render_pptx_slides',return_value={1:'pages/test_s001.png'}):
                with self.assertRaisesRegex(RuntimeError,'slide count'):
                    ingest._ingest_pptx(str(deck))

    def test_render_uses_isolated_profile_and_preserves_hidden_slides(self):
        import fitz
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            images=root/'pages'; images.mkdir()
            deck=root/'source.pptx'; deck.touch()
            calls=[]
            def convert(args, **kwargs):
                calls.append((args,kwargs))
                out=Path(args[args.index('--outdir')+1])/'source.pdf'
                with fitz.open() as pdf:
                    pdf.new_page(); pdf.new_page(); pdf.save(out)
            with patch.object(ingest.config,'DATA_INDEX',td), patch.object(ingest,'_find_soffice',return_value='/opt/soffice'), patch.object(ingest.subprocess,'run',side_effect=convert):
                result=ingest._render_pptx_slides(str(deck),'source',str(images))
            args,kwargs=calls[0]
            self.assertTrue(any(a.startswith('-env:UserInstallation=') for a in args))
            self.assertIn('ExportHiddenSlides',args[args.index('--convert-to')+1])
            self.assertEqual(kwargs.get('timeout'),120)
            self.assertEqual(set(result),{1,2})
            self.assertTrue(all((root/value).is_file() for value in result.values()))


if __name__ == '__main__':
    unittest.main()
