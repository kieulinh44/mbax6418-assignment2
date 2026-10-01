"""Native-resolution visual evidence from the exact retrieved PPTX slide."""
import base64
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import llm, main


def png(size, color):
    stream = io.BytesIO()
    Image.new("RGB", size, color).save(stream, format="PNG")
    stream.seek(0)
    return stream


def decode(uri):
    assert uri.startswith("data:image/png;base64,")
    return Image.open(io.BytesIO(base64.b64decode(uri.split(",", 1)[1])))


class VisualDetailImagesTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.raw = Path(temp.name)
        patcher = patch.object(main.config, "DATA_RAW", str(self.raw))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.deck = Presentation()
        self.deck.slide_width = Inches(10)
        self.deck.slide_height = Inches(5.625)
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        other = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        other.shapes.add_picture(png((1024, 512), "blue"), Inches(1), Inches(1),
                                 width=Inches(6), height=Inches(3))
        self.full = str(self.raw / "full-slide.png")
        self.source = {"file": "lecture.pptx", "page": 1, "kind": "pptx"}

    def save(self):
        self.deck.save(self.raw / "lecture.pptx")

    def test_exact_slide_native_visible_crop_excludes_logos_and_other_slide(self):
        # A real 1024x512 original occupying a substantial fraction of the slide.
        original = Image.new("RGB", (1024, 512), "red")
        original.paste("green", (256, 128, 768, 384))
        stream = io.BytesIO()
        original.save(stream, format="PNG")
        stream.seek(0)
        picture = self.slide.shapes.add_picture(stream, Inches(1), Inches(1),
                                                width=Inches(6), height=Inches(3))
        picture.crop_left = picture.crop_right = 0.25
        picture.crop_top = picture.crop_bottom = 0.25
        self.slide.shapes.add_picture(png((1024, 512), "yellow"), 0, 0,
                                     width=Inches(0.5), height=Inches(0.25))
        self.slide.shapes.add_picture(png((127, 512), "purple"), 0, 0,
                                     width=Inches(4), height=Inches(3))
        self.save()
        images = main._slide_vision_images(self.source, self.full)
        self.assertEqual(images[0], self.full)
        self.assertEqual(len(images), 2)
        crop = decode(images[1])
        self.assertEqual(crop.size, (512, 256))
        self.assertEqual(crop.getpixel((0, 0)), (0, 128, 0))
        self.assertEqual(crop.getpixel((511, 255)), (0, 128, 0))
        self.assertEqual(crop.getextrema(), ((0, 0), (128, 128), (0, 0)))
        self.assertEqual(sorted(p.name for p in self.raw.iterdir()), ["lecture.pptx"])

    def test_no_more_than_three_details_at_native_resolution(self):
        for color in ("red", "green", "yellow", "purple"):
            self.slide.shapes.add_picture(png((1024, 512), color), Inches(1), Inches(1),
                                         width=Inches(6), height=Inches(3))
        self.save()
        images = main._slide_vision_images(self.source, self.full)
        self.assertEqual(len(images), 4)
        self.assertEqual(images[0], self.full)
        self.assertEqual([decode(uri).size for uri in images[1:]], [(1024, 512)] * 3)

    def test_full_slide_fallback_without_usable_exact_raw_source(self):
        for source in (self.source, {"file": "reading.pdf", "page": 1},
                       {"file": "lecture.pptx", "page": 0},
                       {"file": "lecture.pptx", "page": 3}):
            with self.subTest(source=source):
                self.assertEqual(main._slide_vision_images(source, self.full), [self.full])
        (self.raw / "lecture.pptx").write_bytes(b"not a pptx")
        self.assertEqual(main._slide_vision_images(self.source, self.full), [self.full])
        self.save()
        self.assertEqual(main._slide_vision_images(self.source, self.full), [self.full])
        for page in (0, 3):
            self.assertEqual(main._slide_vision_images({**self.source, "page": page}, self.full),
                             [self.full])

    def test_raw_lookup_uses_basename_and_one_based_slide(self):
        self.save()
        images = main._slide_vision_images({"file": "../../lecture.pptx", "page": 2}, self.full)
        self.assertEqual(len(images), 2)
        self.assertEqual(decode(images[1]).size, (1024, 512))
        self.assertEqual(decode(images[1]).getpixel((0, 0)), (0, 0, 255))

    def test_ask_sends_two_original_charts_as_one_source(self):
        for color in ("red", "green"):
            self.slide.shapes.add_picture(png((1024, 512), color), Inches(1), Inches(1),
                                         width=Inches(6), height=Inches(3))
        self.save()
        Path(self.full).write_bytes(png((1200, 676), "white").getvalue())
        page = {"doc": "lecture", "file": "lecture.pptx", "page": 1,
                "image": "full-slide.png", "text": "Benchmark charts"}
        with patch.object(main.config, "DATA_INDEX", str(self.raw)), \
                patch.object(main.hybrid, "get_corpus") as corpus, \
                patch.object(main.llm, "chat", return_value="Answer [lecture.pptx p.1]"), \
                patch.object(main.llm, "ground", return_value="Answer [lecture.pptx p.1]"), \
                patch.object(main.llm, "vision", return_value="Two chart observations") as vision:
            corpus.return_value.retrieve.return_value = [{"page": page, "chunk": page["text"]}]
            result = main.ask({"question": "Read the small chart labels"})
        prompt, images = vision.call_args.args
        self.assertEqual(images[0], self.full)
        self.assertEqual(len(images), 3)
        self.assertEqual([decode(uri).size for uri in images[1:]], [(1024, 512)] * 2)
        self.assertIn("exact slide", prompt.lower())
        self.assertIn("not separate sources", prompt.lower())
        self.assertIn("native-resolution", prompt.lower())
        self.assertIn("complete retrieved slide", prompt.lower())
        self.assertIn("unreadable labels", prompt)
        self.assertEqual(result["vision_sources"], [{"file": "lecture.pptx", "page": 1,
                         "status": "success", "notes": "Two chart observations"}])

    def test_vision_passes_internal_data_uri_without_opening_filename(self):
        uri = "data:image/png;base64," + base64.b64encode(png((128, 128), "red").getvalue()).decode()
        with patch.object(llm, "_complete", return_value="Observed chart") as complete:
            with patch("builtins.open", side_effect=AssertionError("URI is not a filename")):
                self.assertEqual(llm.vision("Read the chart", [uri]), "Observed chart")
        content = complete.call_args.args[3][0]["content"]
        self.assertEqual(content[1]["image_url"]["url"], uri)
        self.assertEqual(complete.call_args.kwargs["max_tokens"], 4096)
